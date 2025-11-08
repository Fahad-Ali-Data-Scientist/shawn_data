"""
Export handler for live camera detection results
Supports PDF, DOCX, and Excel formats
"""
import os
from datetime import datetime
from io import BytesIO

def export_to_pdf(session_data, output_path=None):
    """
    Export session data to PDF format
    
    Args:
        session_data: Dictionary containing session information
        output_path: Optional path to save the PDF
    
    Returns:
        BytesIO or file path
    """
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    
    # Container for elements
    elements = []
    styles = getSampleStyleSheet()
    
    # Title style
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#2563eb'),
        spaceAfter=30,
        alignment=TA_CENTER
    )
    
    # Add title
    elements.append(Paragraph("Live Camera Detection Report", title_style))
    elements.append(Spacer(1, 0.2 * inch))
    
    # Session Information
    session_info = [
        ['Session ID:', session_data.get('session_id', 'N/A')],
        ['Generated:', datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
        ['Total Frames Processed:', str(session_data.get('total_frames', 0))],
        ['Total Objects Detected:', str(session_data.get('total_objects_detected', 0))],
        ['Session Duration:', f"{session_data.get('duration', 0):.2f} seconds"]
    ]
    
    info_table = Table(session_info, colWidths=[2*inch, 4*inch])
    info_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f1f5f9')),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
        ('TOPPADDING', (0, 0), (-1, -1), 12),
        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#e2e8f0'))
    ]))
    
    elements.append(info_table)
    elements.append(Spacer(1, 0.3 * inch))
    
    # Object Detection Summary
    elements.append(Paragraph("Object Detection Summary", styles['Heading2']))
    elements.append(Spacer(1, 0.1 * inch))
    
    class_counts = session_data.get('class_counts', {})
    if class_counts:
        detection_data = [['Object Class', 'Total Count', 'Percentage']]
        total_detections = sum(class_counts.values())
        
        for class_name, count in sorted(class_counts.items(), key=lambda x: x[1], reverse=True):
            percentage = (count / total_detections * 100) if total_detections > 0 else 0
            detection_data.append([class_name, str(count), f"{percentage:.1f}%"])
        
        detection_table = Table(detection_data, colWidths=[2.5*inch, 1.5*inch, 1.5*inch])
        detection_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563eb')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('TOPPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')])
        ]))
        
        elements.append(detection_table)
    else:
        elements.append(Paragraph("No objects detected during this session.", styles['Normal']))
    
    # Build PDF
    doc.build(elements)
    
    if output_path:
        with open(output_path, 'wb') as f:
            f.write(buffer.getvalue())
        return output_path
    
    buffer.seek(0)
    return buffer

def export_to_docx(session_data, output_path=None):
    """
    Export session data to DOCX format
    
    Args:
        session_data: Dictionary containing session information
        output_path: Optional path to save the DOCX
    
    Returns:
        BytesIO or file path
    """
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    
    doc = Document()
    
    # Add title
    title = doc.add_heading('Live Camera Detection Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.runs[0]
    title_run.font.color.rgb = RGBColor(37, 99, 235)
    
    doc.add_paragraph()
    
    # Session Information
    doc.add_heading('Session Information', level=1)
    
    info_table = doc.add_table(rows=5, cols=2)
    info_table.style = 'Light Grid Accent 1'
    
    info_data = [
        ('Session ID', session_data.get('session_id', 'N/A')),
        ('Generated', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
        ('Total Frames Processed', str(session_data.get('total_frames', 0))),
        ('Total Objects Detected', str(session_data.get('total_objects_detected', 0))),
        ('Session Duration', f"{session_data.get('duration', 0):.2f} seconds")
    ]
    
    for i, (label, value) in enumerate(info_data):
        info_table.rows[i].cells[0].text = label
        info_table.rows[i].cells[1].text = value
        # Make labels bold
        info_table.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True
    
    doc.add_paragraph()
    
    # Object Detection Summary
    doc.add_heading('Object Detection Summary', level=1)
    
    class_counts = session_data.get('class_counts', {})
    if class_counts:
        detection_table = doc.add_table(rows=len(class_counts) + 1, cols=3)
        detection_table.style = 'Light Grid Accent 1'
        
        # Header
        header_cells = detection_table.rows[0].cells
        header_cells[0].text = 'Object Class'
        header_cells[1].text = 'Total Count'
        header_cells[2].text = 'Percentage'
        
        for cell in header_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
        
        # Data
        total_detections = sum(class_counts.values())
        sorted_classes = sorted(class_counts.items(), key=lambda x: x[1], reverse=True)
        
        for i, (class_name, count) in enumerate(sorted_classes, start=1):
            row_cells = detection_table.rows[i].cells
            percentage = (count / total_detections * 100) if total_detections > 0 else 0
            row_cells[0].text = class_name
            row_cells[1].text = str(count)
            row_cells[2].text = f"{percentage:.1f}%"
    else:
        doc.add_paragraph('No objects detected during this session.')
    
    # Footer
    doc.add_paragraph()
    footer_para = doc.add_paragraph('Generated by YOLO Detection System')
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.runs[0]
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(100, 116, 139)
    
    if output_path:
        doc.save(output_path)
        return output_path
    
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def export_to_excel(session_data, output_path=None):
    """
    Export session data to Excel format
    
    Args:
        session_data: Dictionary containing session information
        output_path: Optional path to save the Excel file
    
    Returns:
        BytesIO or file path
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    
    wb = Workbook()
    
    # Summary Sheet
    ws1 = wb.active
    ws1.title = "Session Summary"
    
    # Header styling
    header_fill = PatternFill(start_color="2563eb", end_color="2563eb", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    
    # Title
    ws1['A1'] = 'Live Camera Detection Report'
    ws1['A1'].font = Font(size=16, bold=True, color="2563eb")
    ws1.merge_cells('A1:B1')
    
    # Session Info
    ws1['A3'] = 'Session Information'
    ws1['A3'].font = Font(size=14, bold=True)
    
    info_data = [
        ('Session ID', session_data.get('session_id', 'N/A')),
        ('Generated', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
        ('Total Frames Processed', session_data.get('total_frames', 0)),
        ('Total Objects Detected', session_data.get('total_objects_detected', 0)),
        ('Session Duration (seconds)', f"{session_data.get('duration', 0):.2f}")
    ]
    
    row = 4
    for label, value in info_data:
        ws1[f'A{row}'] = label
        ws1[f'B{row}'] = value
        ws1[f'A{row}'].font = Font(bold=True)
        row += 1
    
    # Detection Summary Sheet
    ws2 = wb.create_sheet("Detection Summary")
    
    # Headers
    ws2['A1'] = 'Object Class'
    ws2['B1'] = 'Total Count'
    ws2['C1'] = 'Percentage'
    
    for cell in ['A1', 'B1', 'C1']:
        ws2[cell].fill = header_fill
        ws2[cell].font = header_font
        ws2[cell].alignment = Alignment(horizontal='center')
    
    # Data
    class_counts = session_data.get('class_counts', {})
    total_detections = sum(class_counts.values())
    sorted_classes = sorted(class_counts.items(), key=lambda x: x[1], reverse=True)
    
    row = 2
    for class_name, count in sorted_classes:
        percentage = (count / total_detections * 100) if total_detections > 0 else 0
        ws2[f'A{row}'] = class_name
        ws2[f'B{row}'] = count
        ws2[f'C{row}'] = f"{percentage:.1f}%"
        row += 1
    
    # Auto-adjust column widths
    for ws in [ws1, ws2]:
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    # Detection History Sheet (if available)
    detections_history = session_data.get('detections_history', [])
    if detections_history:
        ws3 = wb.create_sheet("Detection History")
        
        ws3['A1'] = 'Frame'
        ws3['B1'] = 'Timestamp'
        ws3['C1'] = 'Objects Detected'
        ws3['D1'] = 'Object Classes'
        
        for cell in ['A1', 'B1', 'C1', 'D1']:
            ws3[cell].fill = header_fill
            ws3[cell].font = header_font
        
        row = 2
        for entry in detections_history:
            ws3[f'A{row}'] = entry.get('frame', 0)
            ws3[f'B{row}'] = entry.get('timestamp', '')
            ws3[f'C{row}'] = entry.get('count', 0)
            
            # Get unique classes in this frame
            detections = entry.get('detections', [])
            classes = ', '.join(set([d['class'] for d in detections]))
            ws3[f'D{row}'] = classes
            row += 1
        
        # Auto-adjust columns
        for column in ws3.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws3.column_dimensions[column_letter].width = adjusted_width
    
    if output_path:
        wb.save(output_path)
        return output_path
    
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

